# Извлечение текста из презентации и реферата (НГАСУ).
import sys
from pptx import Presentation
from docx import Document

BASE = r'C:\Users\user\Desktop\Чат-боты (по докладу НГАСУ)'

print('=' * 30, 'PPTX', '=' * 30)
prs = Presentation(f'{BASE}\\Презентация_Чат-боты_архитектура_строительство.pptx')


def walk(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # GROUP
            walk(shape.shapes)
            continue
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = ''.join(run.text for run in para.runs).strip()
                if text:
                    print(text)
        if shape.has_table:
            for row in shape.table.rows:
                cells = [' '.join(c.text.split()) for c in row.cells]
                print(' | '.join(cells))


for i, slide in enumerate(prs.slides, 1):
    print(f'\n--- СЛАЙД {i} ---')
    walk(slide.shapes)

print('\n', '=' * 30, 'DOCX', '=' * 30)
doc = Document(f'{BASE}\\Реферат_Чат-боты_архитектура_строительство.docx')
for para in doc.paragraphs:
    text = para.text.strip()
    if text:
        print(text)
